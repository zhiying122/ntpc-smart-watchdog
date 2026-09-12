# -*- coding: utf-8 -*-
"""
小小守護員 Smart Watchdog — 上台報告 PPT 產生器
================================================
依實際系統成果（src/forensic.py 鑑識會計、src/risk_score.py 白盒分數、
src/ai_report.py Bedrock、雙 app 8601/8602、AWS 架構）產出決賽簡報。

對齊評審必備四項：解決方案說明、數據及資料運用、AWS 雲端技術架構(架構圖)、
解決方案使用介面及操作流程。

用法：
    python scripts/build_pitch_ppt.py
輸出：
    reports/小小守護員_決賽簡報.pptx
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "reports", "小小守護員_決賽簡報.pptx")

# 16:9
EMU = 914400
SW = int(13.333 * EMU)
SH = int(7.5 * EMU)

# ---- 政府/企業級中性色盤（對齊 ui-design steering）----
INK = RGBColor(0x1A, 0x1F, 0x2B)        # 深墨（主文字）
INK_MUTED = RGBColor(0x5B, 0x63, 0x72)  # 次要文字
BG = RGBColor(0xFF, 0xFF, 0xFF)         # 白
SURFACE = RGBColor(0xF4, 0xF6, 0xF9)    # 淺灰底
SURFACE2 = RGBColor(0xEA, 0xEE, 0xF3)
BORDER = RGBColor(0xD5, 0xDB, 0xE3)
PRIMARY = RGBColor(0x1F, 0x4E, 0x79)    # 政府藍（主互動/資訊）
PRIMARY_DK = RGBColor(0x14, 0x33, 0x50)
ACCENT = RGBColor(0x2C, 0x6E, 0x9B)
# 風險五階語意色
RISK_RED = RGBColor(0xD6, 0x45, 0x45)
RISK_ORANGE = RGBColor(0xE0, 0x7B, 0x2B)
RISK_YELLOW = RGBColor(0xE0, 0xA9, 0x3B)
RISK_GREEN = RGBColor(0x3F, 0x8F, 0x5B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Microsoft JhengHei"  # Noto Sans TC 替代（Windows 中文黑體）

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- helpers
def slide():
    return prs.slides.add_slide(BLANK)


def _set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(s, x, y, w, h, color, line=None, line_w=None, shadow=False,
         rounded=False):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Emu(x), Emu(y), Emu(w), Emu(h))
    if color is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = color
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w or 1)
    shp.shadow.inherit = False
    return shp


def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space_after=4, line_spacing=1.0, wrap=True):
    """runs: list of (text, size, bold, color) OR list of paragraphs where
    each paragraph is a list of such run-tuples."""
    tb = s.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    # normalise: list of paragraphs
    if runs and isinstance(runs[0], tuple):
        paras = [runs]
    else:
        paras = runs

    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (t, size, bold, color) in para:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = FONT
            r.font.color.rgb = color
    return tb


def page_header(s, kicker, title, idx=None):
    """統一頁首：左上細色塊 + kicker + 大標題；右上頁碼。"""
    rect(s, 0, 0, SW, Emu(int(1.35 * EMU)), BG)
    rect(s, int(0.6 * EMU), int(0.42 * EMU), int(0.09 * EMU),
         int(0.62 * EMU), PRIMARY)
    txt(s, int(0.85 * EMU), int(0.38 * EMU), int(10 * EMU), int(0.3 * EMU),
        [(kicker, 12, True, ACCENT)])
    txt(s, int(0.85 * EMU), int(0.62 * EMU), int(11.2 * EMU), int(0.62 * EMU),
        [(title, 27, True, INK)])
    # 底部細線
    rect(s, int(0.85 * EMU), int(1.24 * EMU), int(11.6 * EMU), Emu(12700),
         BORDER)
    if idx is not None:
        txt(s, int(11.9 * EMU), int(0.5 * EMU), int(0.9 * EMU),
            int(0.3 * EMU), [(f"{idx:02d}", 12, True, BORDER)],
            align=PP_ALIGN.RIGHT)


def chip(s, x, y, w, h, label, fill, fg=WHITE, size=11, bold=True):
    c = rect(s, x, y, w, h, fill, rounded=True)
    tf = c.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(4)
    tf.margin_right = Pt(4)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = FONT
    r.font.color.rgb = fg
    return c


def card(s, x, y, w, h, fill=SURFACE, line=BORDER):
    return rect(s, x, y, w, h, fill, line=line, line_w=1, rounded=True)


IN = EMU  # inch shorthand


# ================================================================ SLIDE 1 封面
def s_cover():
    s = slide()
    rect(s, 0, 0, SW, SH, PRIMARY_DK)
    # 左側主色帶
    rect(s, 0, 0, int(0.16 * IN), SH, PRIMARY)
    # 頂部標籤
    chip(s, int(0.85 * IN), int(0.9 * IN), int(5.7 * IN), int(0.42 * IN),
         "2026 新北市 AI 智慧城市黑客松競賽 ｜ 教育局組",
         RGBColor(0x2A, 0x4A, 0x6B), fg=RGBColor(0xCF, 0xE0, 0xF0), size=13)

    txt(s, int(0.85 * IN), int(2.0 * IN), int(11.5 * IN), int(1.6 * IN),
        [[("小小守護員 ", 52, True, WHITE),
          ("Smart Watchdog", 52, True, RGBColor(0x7F, 0xB5, 0xDA))]],
        space_after=2)
    txt(s, int(0.85 * IN), int(3.35 * IN), int(11.5 * IN), int(0.8 * IN),
        [("AI × 鑑識會計，打造教保機構智慧風險預警管理系統", 22, True,
          RGBColor(0xE6, 0xEE, 0xF6))])

    # 一句話定位條
    bar = rect(s, int(0.85 * IN), int(4.5 * IN), int(11.5 * IN),
               int(1.0 * IN), RGBColor(0x24, 0x42, 0x60), rounded=True)
    txt(s, int(1.15 * IN), int(4.6 * IN), int(11.0 * IN), int(0.8 * IN),
        [[("別人做幼兒園風險評分，我們做「", 15, False,
           RGBColor(0xD7, 0xE3, 0xEF)),
          ("AI 稽查官", 15, True, RGBColor(0xFF, 0xD9, 0x7A)),
          ("」──用審計界抓弊案的鑑識會計算出", 15, False,
           RGBColor(0xD7, 0xE3, 0xEF))],
         [("可解釋的白盒子風險分數，再讓生成式 AI 產出稽查建議，讓稽查人力產能放大十倍。",
           15, False, RGBColor(0xD7, 0xE3, 0xEF))]],
        anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.15)

    # 底部：雙系統網址
    txt(s, int(0.85 * IN), int(6.15 * IN), int(11.5 * IN), int(0.9 * IN),
        [[("公務後台（稽查官）", 12, True, RGBColor(0x9F, 0xBE, 0xDC)),
          ("  :8601        ", 12, False, RGBColor(0x7F, 0x9C, 0xBA)),
          ("公眾查詢網（家長）", 12, True, RGBColor(0x9F, 0xBE, 0xDC)),
          ("  :8602", 12, False, RGBColor(0x7F, 0x9C, 0xBA))],
         [("Live on AWS EC2 · systemd 常駐 · GitHub Actions 自動部署",
           11, False, RGBColor(0x6E, 0x8A, 0xA6))]],
        space_after=3)


# ================================================================ SLIDE 2 痛點
def s_pain():
    s = slide()
    page_header(s, "BACKGROUND ｜ 背景與痛點", "現行監理的四個結構性缺口", 2)
    intro = ("近年教保機構違法及不當管理事件時有發生，衝擊幼兒受教權益與政府監督公信力。"
             "現行機制仰賴定期稽查、書面審核與被動通報，難以事前辨識高風險機構。")
    txt(s, int(0.85 * IN), int(1.45 * IN), int(11.6 * IN), int(0.6 * IN),
        [(intro, 13, False, INK_MUTED)], line_spacing=1.15)

    pains = [
        ("01", "資料分散、整合困難",
         "基本資料、評鑑、裁罰、收費、決算雖為公開資訊，卻散落各處、格式不一，"
         "且未納入社群即時輿情，無法形成完整風險判斷。"),
        ("02", "財務異常難及早發現",
         "決算與財報雖依法公告，卻缺乏與收費、園所規模的交叉勾稽，"
         "數字造假與經費異常往往事後才被發現。"),
        ("03", "缺乏可量化的預警工具",
         "沒有標準化、可量化的風險評估方法，難以從大量機構中篩出"
         "「該優先關注的高風險對象」。"),
        ("04", "人力吃緊、決策缺依據",
         "機構數量龐大，仰賴人工逐一篩選耗時費力，"
         "資源配置缺乏數據支撐的決策依據。"),
    ]
    x0 = int(0.85 * IN)
    gap = int(0.3 * IN)
    cw = (int(11.6 * IN) - gap) // 2
    ch = int(1.85 * IN)
    y0 = int(2.15 * IN)
    for i, (num, title, body) in enumerate(pains):
        col = i % 2
        row = i // 2
        x = x0 + col * (cw + gap)
        y = y0 + row * (ch + int(0.25 * IN))
        card(s, x, y, cw, ch)
        rect(s, x, y, int(0.09 * IN), ch, RISK_ORANGE)
        txt(s, x + int(0.28 * IN), y + int(0.2 * IN), int(1.0 * IN),
            int(0.5 * IN), [(num, 26, True, RGBColor(0xC6, 0xCE, 0xD8))])
        txt(s, x + int(1.15 * IN), y + int(0.24 * IN), cw - int(1.4 * IN),
            int(0.4 * IN), [(title, 16, True, INK)])
        txt(s, x + int(1.15 * IN), y + int(0.72 * IN), cw - int(1.4 * IN),
            int(1.0 * IN), [(body, 11.5, False, INK_MUTED)], line_spacing=1.15)


# ================================================================ SLIDE 3 解決方案總覽
def s_solution():
    s = slide()
    page_header(s, "SOLUTION ｜ 解決方案說明", "抓資料 → 算風險 → 出建議：三段式 AI 稽查官", 3)

    steps = [
        ("抓資料", "整合公開多源資料",
         ["公校決算書（112–114 年度，文字層抽取）",
          "非營利園掃描財報（Bedrock 多模態 / OCR）",
          "裁罰、評鑑、收費、社群輿情"],
         PRIMARY),
        ("算風險", "鑑識會計白盒子分數",
         ["班佛定律、Beneish、Altman、孤立森林",
          "財務比率交叉勾稽、基金餘額勾稽",
          "0–100 可攤開解釋，非黑盒子"],
         ACCENT),
        ("出建議", "生成式 AI 稽查輔助",
         ["AWS Bedrock（Claude）翻譯數字成白話",
          "指出異常、給具體查核方向",
          "責任 AI：風險 ≠ 違法，附證據鏈"],
         RISK_GREEN),
    ]
    x0 = int(0.85 * IN)
    gap = int(0.35 * IN)
    cw = (int(11.6 * IN) - 2 * gap) // 3
    y0 = int(1.7 * IN)
    ch = int(3.5 * IN)
    for i, (tag, sub, items, color) in enumerate(steps):
        x = x0 + i * (cw + gap)
        card(s, x, y0, cw, ch, fill=WHITE)
        rect(s, x, y0, cw, int(0.85 * IN), color, rounded=True)
        rect(s, x, y0 + int(0.6 * IN), cw, int(0.25 * IN), color)
        chip(s, x + int(0.25 * IN), y0 + int(0.22 * IN), int(0.5 * IN),
             int(0.42 * IN), str(i + 1), WHITE, fg=color, size=17)
        txt(s, x + int(0.9 * IN), y0 + int(0.16 * IN), cw - int(1.0 * IN),
            int(0.6 * IN), [(tag, 20, True, WHITE)],
            anchor=MSO_ANCHOR.MIDDLE)
        txt(s, x + int(0.3 * IN), y0 + int(1.05 * IN), cw - int(0.6 * IN),
            int(0.4 * IN), [(sub, 14, True, INK)])
        yy = y0 + int(1.6 * IN)
        for it in items:
            rect(s, x + int(0.32 * IN), yy + int(0.08 * IN), int(0.12 * IN),
                 int(0.12 * IN), color)
            txt(s, x + int(0.55 * IN), yy, cw - int(0.85 * IN),
                int(0.6 * IN), [(it, 11.5, False, INK_MUTED)],
                line_spacing=1.1)
            yy += int(0.62 * IN)

    # 底部效益條
    y = int(5.5 * IN)
    benefits = [("提升風險辨識率", "從公開資料篩出高風險機構"),
                ("縮短預警時間", "事後被動稽查 → 事前主動示警"),
                ("降低人力負擔", "有限人力精準投放最該查的園")]
    bw = (int(11.6 * IN) - 2 * int(0.3 * IN)) // 3
    for i, (t, d) in enumerate(benefits):
        x = x0 + i * (bw + int(0.3 * IN))
        rect(s, x, y, bw, int(1.15 * IN), SURFACE, line=BORDER, line_w=1,
             rounded=True)
        rect(s, x, y, int(0.09 * IN), int(1.15 * IN), RISK_GREEN)
        txt(s, x + int(0.28 * IN), y + int(0.18 * IN), bw - int(0.4 * IN),
            int(0.4 * IN), [(t, 14, True, INK)])
        txt(s, x + int(0.28 * IN), y + int(0.62 * IN), bw - int(0.4 * IN),
            int(0.45 * IN), [(d, 11, False, INK_MUTED)], line_spacing=1.1)


# ================================================================ SLIDE 4 護城河：鑑識會計
def s_forensic():
    s = slide()
    page_header(s, "CORE ｜ 差異化護城河", "鑑識會計搬進幼兒園：多層方法，全部可解釋", 4)
    txt(s, int(0.85 * IN), int(1.45 * IN), int(11.6 * IN), int(0.5 * IN),
        [[("90% 隊伍只做「風險評分」，", 13, False, INK_MUTED),
          ("我們用審計界抓弊案的鑑識會計方法", 13, True, PRIMARY),
          ("，每一層都有學術文獻佐證。", 13, False, INK_MUTED)]],
        line_spacing=1.15)

    methods = [
        ("班佛定律 Benford's Law", "逐筆明細金額首位數分布偏離自然律 → 數字造假傾向（MAD + 卡方檢定）", "Nigrini；Durtschi 2004"),
        ("Beneish M-Score 改良版", "收入與支出成長背離、應計項目異常 → 盈餘操縱（在地化為非營利版）", "Beneish 1999"),
        ("Altman Z'' 困境分", "財務困境代理比率 → 對應舞弊三角的「壓力」構面", "Altman 1968 / 2013"),
        ("Isolation Forest + SHAP", "非監督式多維異常偵測，SHAP 解釋「為何異常」，維持白盒", "Protiviti 白皮書"),
        ("財務比率交叉勾稽 + IQR", "每幼兒單位成本、人事費占比、收支比 vs 同儕群體離群", "PCAOB AS 2305"),
        ("基金餘額勾稽", "期末 = 期初 + 本期賸餘；跨年度連續性核對，接不上即紅旗", "政府基金會計恆等式"),
    ]
    x0 = int(0.85 * IN)
    gap = int(0.28 * IN)
    cw = (int(11.6 * IN) - 2 * gap) // 3
    ch = int(1.7 * IN)
    y0 = int(2.15 * IN)
    for i, (name, desc, ref) in enumerate(methods):
        col = i % 3
        row = i // 3
        x = x0 + col * (cw + gap)
        y = y0 + row * (ch + int(0.25 * IN))
        card(s, x, y, cw, ch, fill=WHITE)
        rect(s, x, y, cw, int(0.08 * IN), PRIMARY)
        txt(s, x + int(0.24 * IN), y + int(0.2 * IN), cw - int(0.45 * IN),
            int(0.55 * IN), [(name, 13.5, True, INK)], line_spacing=1.0)
        txt(s, x + int(0.24 * IN), y + int(0.78 * IN), cw - int(0.45 * IN),
            int(0.7 * IN), [(desc, 10.5, False, INK_MUTED)], line_spacing=1.1)
        txt(s, x + int(0.24 * IN), y + int(1.42 * IN), cw - int(0.45 * IN),
            int(0.25 * IN),
            [[("文獻：", 9, True, ACCENT), (ref, 9, False, INK_MUTED)]])


# ================================================================ SLIDE 5 白盒子分數
def s_whitebox():
    s = slide()
    page_header(s, "EXPLAINABLE ｜ 白盒子風險分數", "禁止黑盒子：每一分都攤得開", 5)

    # 左：權重公式卡
    x0 = int(0.85 * IN)
    lw = int(5.5 * IN)
    card(s, x0, int(1.6 * IN), lw, int(2.4 * IN), fill=SURFACE)
    txt(s, x0 + int(0.3 * IN), int(1.8 * IN), lw - int(0.6 * IN),
        int(0.4 * IN), [("分項權重（可解釋加權）", 15, True, INK)])
    weights = [("財務異常", "50%", "鑑識會計多層方法", RISK_RED),
               ("裁罰紀錄", "34%", "事由分類 + 嚴重度 / 破窗效應", RISK_ORANGE),
               ("評鑑結果", "16%", "優 0 / 良 20 / 乙 40 / 待改進 90", RISK_YELLOW)]
    yy = int(2.32 * IN)
    for name, pct, desc, color in weights:
        rect(s, x0 + int(0.3 * IN), yy + int(0.05 * IN), int(0.14 * IN),
             int(0.5 * IN), color)
        txt(s, x0 + int(0.55 * IN), yy, int(1.4 * IN), int(0.5 * IN),
            [[(name + "  ", 13, True, INK), (pct, 15, True, color)]])
        txt(s, x0 + int(2.2 * IN), yy + int(0.06 * IN), int(3.1 * IN),
            int(0.5 * IN), [(desc, 10, False, INK_MUTED)], line_spacing=1.0)
        yy += int(0.56 * IN)
    txt(s, x0 + int(0.3 * IN), int(3.72 * IN), lw - int(0.6 * IN),
        int(0.3 * IN),
        [("＊輿情分析已實作並於介面呈現；總分計分待真實爬蟲全量接入後啟用（誠實揭露）",
          9, False, INK_MUTED)], line_spacing=1.0)

    # 右：實例拆解（林口幼兒園）
    rx = x0 + lw + int(0.4 * IN)
    rw = int(11.6 * IN) - lw - int(0.4 * IN)
    card(s, rx, int(1.6 * IN), rw, int(2.4 * IN), fill=WHITE)
    rect(s, rx, int(1.6 * IN), rw, int(0.55 * IN), PRIMARY, rounded=True)
    rect(s, rx, int(1.95 * IN), rw, int(0.2 * IN), PRIMARY)
    txt(s, rx + int(0.3 * IN), int(1.68 * IN), rw - int(0.6 * IN),
        int(0.4 * IN), [("判定範例：新北市立林口幼兒園 52.6 分", 14, True, WHITE)],
        anchor=MSO_ANCHOR.MIDDLE)
    lines = [("財務異常  38.0 × 0.50", "= 19.0", RISK_RED),
             ("裁罰(2次) 80  × 0.34", "= 27.2", RISK_ORANGE),
             ("評鑑(乙)  40  × 0.16", "=  6.4", RISK_YELLOW)]
    yy = int(2.35 * IN)
    for left, right, color in lines:
        txt(s, rx + int(0.35 * IN), yy, int(3.6 * IN), int(0.4 * IN),
            [(left, 13, False, INK)])
        txt(s, rx + int(3.7 * IN), yy, int(1.4 * IN), int(0.4 * IN),
            [(right, 13, True, color)], align=PP_ALIGN.RIGHT)
        yy += int(0.5 * IN)
    rect(s, rx + int(0.35 * IN), yy, rw - int(0.7 * IN), Emu(12700), BORDER)
    yy += int(0.12 * IN)
    txt(s, rx + int(0.35 * IN), yy, int(3.6 * IN), int(0.4 * IN),
        [("總風險分（相對分級：高）", 13, True, INK)])
    txt(s, rx + int(3.7 * IN), yy, int(1.4 * IN), int(0.4 * IN),
        [("52.6", 17, True, RISK_RED)], align=PP_ALIGN.RIGHT)

    # 底：模型驗證數字
    y = int(4.35 * IN)
    txt(s, x0, y, int(11.6 * IN), int(0.35 * IN),
        [("模型鑑別力驗證（以官方裁罰為標籤，python -m src.validate 即時產生）",
          13, True, INK)])
    metrics = [("AUC-ROC", "0.906", "辨識力優異 (>0.9)"),
               ("Recall", "1.00", "被裁罰園零漏抓"),
               ("相關係數", "0.733", "風險分與裁罰高度正相關"),
               ("權重敏感度", "ρ≥0.97", "改權重排名幾乎不變")]
    mw = (int(11.6 * IN) - 3 * int(0.25 * IN)) // 4
    yb = int(4.8 * IN)
    for i, (t, v, d) in enumerate(metrics):
        x = x0 + i * (mw + int(0.25 * IN))
        card(s, x, yb, mw, int(1.5 * IN), fill=SURFACE)
        rect(s, x, yb, mw, int(0.07 * IN), RISK_GREEN)
        txt(s, x, yb + int(0.2 * IN), mw, int(0.35 * IN),
            [(t, 12, True, INK_MUTED)], align=PP_ALIGN.CENTER)
        txt(s, x, yb + int(0.55 * IN), mw, int(0.5 * IN),
            [(v, 26, True, PRIMARY)], align=PP_ALIGN.CENTER)
        txt(s, x + int(0.1 * IN), yb + int(1.08 * IN), mw - int(0.2 * IN),
            int(0.4 * IN), [(d, 9.5, False, INK_MUTED)],
            align=PP_ALIGN.CENTER, line_spacing=1.0)


# ================================================================ SLIDE 6 資料運用
def s_data():
    s = slide()
    page_header(s, "DATA ｜ 數據及資料運用", "真實多源整合，誠實揭露抽樣深度與全市涵蓋", 6)

    # 左：資料筆數 KPI
    x0 = int(0.85 * IN)
    kpis = [("1,114", "間", "全市立案機構納管基本資料", PRIMARY),
            ("61", "間", "已完成鑑識會計風險評分", RISK_GREEN),
            ("2,838", "筆", "真實爬取裁罰紀錄快照", RISK_ORANGE),
            ("7,688", "筆", "全國教保機構名冊快照", ACCENT)]
    kw = (int(11.6 * IN) - 3 * int(0.28 * IN)) // 4
    yk = int(1.6 * IN)
    for i, (v, unit, d, color) in enumerate(kpis):
        x = x0 + i * (kw + int(0.28 * IN))
        card(s, x, yk, kw, int(1.5 * IN), fill=WHITE)
        rect(s, x, yk, int(0.09 * IN), int(1.5 * IN), color)
        txt(s, x + int(0.25 * IN), yk + int(0.22 * IN), kw - int(0.4 * IN),
            int(0.55 * IN),
            [[(v, 30, True, color), ("  " + unit, 12, True, INK_MUTED)]])
        txt(s, x + int(0.25 * IN), yk + int(0.92 * IN), kw - int(0.4 * IN),
            int(0.5 * IN), [(d, 11, False, INK_MUTED)], line_spacing=1.1)

    # 資料來源矩陣（誠實三態）
    txt(s, x0, int(3.35 * IN), int(11.6 * IN), int(0.35 * IN),
        [("資料來源矩陣（誠實清冊，src/data_source_matrix.py）", 14, True, INK)])
    rows = [
        ("公校決算書 112–114", "新北市教育局", "PDF 文字層", "已確認", RISK_GREEN, "財務鑑識黃金資料源"),
        ("非營利園財報 110–113", "新北市教育局", "掃描影像", "已確認", RISK_GREEN, "Bedrock 多模態 / OCR 抽取"),
        ("裁罰紀錄（爬取快照）", "全國教保資訊網", "JSON", "已確認", RISK_GREEN, "2,838 筆佐證圖層"),
        ("官方裁罰/評鑑/收費 API", "全國教保資訊網", "未確認", "未確認", INK_MUTED, "接入即套用同流程，架構可規模化"),
        ("網路輿情（NLP 樣本）", "公開新聞/社群", "未確認", "未確認", INK_MUTED, "小樣本示範，暫不計分"),
    ]
    ty = int(3.8 * IN)
    tw = int(11.6 * IN)
    # 表頭
    heads = [("資料集", 3.6), ("主管機關", 2.3), ("格式", 1.6), ("狀態", 1.3), ("用途", 2.8)]
    hx = x0
    rect(s, x0, ty, tw, int(0.42 * IN), PRIMARY_DK)
    for label, wf in heads:
        w = int(wf * IN)
        txt(s, hx + int(0.15 * IN), ty, w - int(0.2 * IN), int(0.42 * IN),
            [(label, 11.5, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
        hx += w
    ry = ty + int(0.42 * IN)
    for ri, (ds, auth, fmt, stat, scol, use) in enumerate(rows):
        rh = int(0.5 * IN)
        rect(s, x0, ry, tw, rh, WHITE if ri % 2 == 0 else SURFACE)
        cells = [(ds, 3.6, INK, True), (auth, 2.3, INK_MUTED, False),
                 (fmt, 1.6, INK_MUTED, False)]
        cx = x0
        for val, wf, col, bold in cells:
            w = int(wf * IN)
            txt(s, cx + int(0.15 * IN), ry, w - int(0.2 * IN), rh,
                [(val, 10.5, bold, col)], anchor=MSO_ANCHOR.MIDDLE)
            cx += w
        # 狀態 chip
        chip(s, cx + int(0.12 * IN), ry + int(0.1 * IN), int(1.0 * IN),
             int(0.3 * IN), stat, scol, size=9.5)
        cx += int(1.3 * IN)
        txt(s, cx + int(0.15 * IN), ry, int(2.8 * IN) - int(0.2 * IN), rh,
            [(use, 10, False, INK_MUTED)], anchor=MSO_ANCHOR.MIDDLE,
            line_spacing=1.0)
        ry += rh
    rect(s, x0, ty, tw, ry - ty, None, line=BORDER, line_w=1)


# ================================================================ SLIDE 7 AWS 架構
def s_aws():
    s = slide()
    page_header(s, "ARCHITECTURE ｜ AWS 雲端技術架構", "資料 → 鑑識引擎 → 雙介面，AWS 全託管", 7)

    x0 = int(0.7 * IN)
    top = int(1.55 * IN)
    full_w = int(11.9 * IN)

    def layer(y, h, title, color, fill):
        rect(s, x0, y, full_w, h, fill, line=BORDER, line_w=1, rounded=True)
        rect(s, x0, y, int(0.09 * IN), h, color)
        txt(s, x0 + int(0.22 * IN), y + int(0.08 * IN), int(3.0 * IN),
            int(0.3 * IN), [(title, 11, True, color)])

    def box(x, y, w, h, title, sub, fill, fg=WHITE, tsize=11.5):
        rect(s, x, y, w, h, fill, rounded=True)
        txt(s, x + int(0.1 * IN), y + int(0.12 * IN), w - int(0.2 * IN),
            int(0.4 * IN), [(title, tsize, True, fg)], align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.95)
        if sub:
            txt(s, x + int(0.08 * IN), y + h - int(0.42 * IN),
                w - int(0.16 * IN), int(0.38 * IN),
                [(sub, 8.5, False, fg)], align=PP_ALIGN.CENTER,
                line_spacing=0.9)

    AWS_ORANGE = RGBColor(0xEC, 0x91, 0x2D)
    # ---- Layer 1 資料來源 ----
    y1 = top
    h1 = int(1.15 * IN)
    layer(y1, h1, "① 資料來源 DATA SOURCES", INK_MUTED, SURFACE)
    srcs = ["公校決算書\n(PDF 文字層)", "非營利園財報\n(掃描影像)",
            "裁罰/評鑑\n(全國教保網)", "網路輿情\n(新聞/社群)"]
    bw = int(2.55 * IN)
    bx = x0 + int(0.9 * IN)
    for t in srcs:
        title, sub = t.split("\n")
        box(bx, y1 + int(0.4 * IN), bw, int(0.62 * IN), title, sub,
            RGBColor(0x5B, 0x63, 0x72))
        bx += bw + int(0.18 * IN)

    # ---- Layer 2 AWS 處理 ----
    y2 = y1 + h1 + int(0.22 * IN)
    h2 = int(1.55 * IN)
    layer(y2, h2, "② AWS 雲端處理層 PROCESSING (boto3)", AWS_ORANGE,
          RGBColor(0xFC, 0xF3, 0xE6))
    aws_boxes = [
        ("Amazon S3", "原始資料 /\n衍生資料儲存"),
        ("AWS Textract\n/ Tesseract OCR", "掃描財報\n文字抽取"),
        ("Amazon Bedrock\nClaude Sonnet 4.5", "多模態抽數字\n+ 生成稽查建議"),
        ("EC2 + IAM Role", "臨時憑證\n免金鑰外洩"),
    ]
    bx = x0 + int(0.9 * IN)
    for title, sub in aws_boxes:
        box(bx, y2 + int(0.42 * IN), bw, int(0.95 * IN), title, sub,
            AWS_ORANGE, tsize=11)
        bx += bw + int(0.18 * IN)

    # ---- Layer 3 鑑識引擎中台 ----
    y3 = y2 + h2 + int(0.22 * IN)
    h3 = int(0.95 * IN)
    layer(y3, h3, "③ 鑑識風險中台 CORE ENGINE (src/)", PRIMARY, SURFACE2)
    eng = ["鑑識會計 forensic.py", "白盒評分 risk_score.py",
           "證據鏈 evidence.py", "AI 報告 ai_report.py", "地理編碼 geocode.py"]
    ew = int(2.0 * IN)
    ex = x0 + int(0.9 * IN)
    for t in eng:
        box(ex, y3 + int(0.32 * IN), ew, int(0.5 * IN), t, "", PRIMARY,
            tsize=9.5)
        ex += ew + int(0.15 * IN)

    # ---- Layer 4 雙介面 ----
    y4 = y3 + h3 + int(0.22 * IN)
    h4 = int(0.85 * IN)
    half = (full_w - int(0.3 * IN)) // 2
    box(x0, y4, half, h4, "公務後台 :8601（稽查官）",
        "風險排名 · 雷達圖 · 案件流 · 派工 · AI 建議", PRIMARY_DK, tsize=13)
    box(x0 + half + int(0.3 * IN), y4, half, h4, "公眾查詢網 :8602（家長）",
        "找幼兒園 · 距離 · 收費 · 公開評鑑 · 輿情觀測", RISK_GREEN, tsize=13)

    # 部署註記
    txt(s, x0, y4 + h4 + int(0.06 * IN), full_w, int(0.3 * IN),
        [[("部署：", 10, True, INK),
          ("GitHub Actions push→main 自動部署到 EC2 self-hosted runner，systemd 常駐兩服務並健康檢查（curl 8601/8602）· 金鑰一律 os.environ，check_secrets.py 上傳前掃描",
           10, False, INK_MUTED)]], line_spacing=1.0)


# ================================================================ SLIDE 8 介面A 公務後台
def s_ui_gov():
    s = slide()
    page_header(s, "UI / FLOW ｜ 介面與操作流程 (1/2)", "公務後台 :8601 — 稽查官的作戰工作台", 8)
    txt(s, int(0.85 * IN), int(1.42 * IN), int(11.6 * IN), int(0.4 * IN),
        [[("使用者：", 12, True, PRIMARY), ("教育局 / 稽查員　｜　", 12, False, INK_MUTED),
          ("動線：全局監控 → 下鑽單園 → 下決定 → 派工結案", 12, True, INK)]])

    flow = [
        ("1  風險總覽", "KPI 帶 + 完整排名表 + 各行政區風險組成堆疊圖，一眼看出該優先查誰"),
        ("2  風險地圖", "Leaflet + OpenStreetMap，紅/黃/綠標記，排稽查路線精準投放"),
        ("3  單園工作台", "雷達圖攤開四分項 + 證據鏈下鑽（分數→特徵→原始決算→PDF 來源）"),
        ("4  AI 稽查建議", "Bedrock 產出白話建議：這間該查什麼憑證，責任 AI 附來源"),
        ("5  案件狀態流", "待研判 → 建議派查 → 調查中 → 結案，附時間戳與稽核軌跡"),
        ("6  派工台 / 資料整合", "排除已稽查、資料來源矩陣、模型鑑別力驗證、系統健康檢查"),
    ]
    x0 = int(0.85 * IN)
    gap = int(0.28 * IN)
    cw = (int(11.6 * IN) - 2 * gap) // 3
    ch = int(1.75 * IN)
    y0 = int(2.05 * IN)
    for i, (t, d) in enumerate(flow):
        col = i % 3
        row = i // 3
        x = x0 + col * (cw + gap)
        y = y0 + row * (ch + int(0.28 * IN))
        card(s, x, y, cw, ch, fill=WHITE)
        rect(s, x, y, cw, int(0.55 * IN), PRIMARY, rounded=True)
        rect(s, x, y + int(0.35 * IN), cw, int(0.2 * IN), PRIMARY)
        txt(s, x + int(0.22 * IN), y + int(0.06 * IN), cw - int(0.4 * IN),
            int(0.45 * IN), [(t, 14, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
        txt(s, x + int(0.24 * IN), y + int(0.72 * IN), cw - int(0.45 * IN),
            int(0.95 * IN), [(d, 11, False, INK_MUTED)], line_spacing=1.15)

    # 案件狀態流 pill 條
    y = int(5.95 * IN)
    states = [("待研判", RISK_YELLOW), ("建議派查", RISK_ORANGE),
              ("調查中", PRIMARY), ("結案·屬實", RISK_RED),
              ("結案·不成立", RISK_GREEN)]
    sx = x0
    sw2 = int(1.9 * IN)
    for i, (name, color) in enumerate(states):
        chip(s, sx, y, sw2, int(0.5 * IN), name, color, size=12)
        if i < len(states) - 1:
            txt(s, sx + sw2, y, int(0.28 * IN), int(0.5 * IN),
                [("→", 16, True, INK_MUTED)], align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
        sx += sw2 + int(0.28 * IN)


# ================================================================ SLIDE 9 介面B 公眾網
def s_ui_public():
    s = slide()
    page_header(s, "UI / FLOW ｜ 介面與操作流程 (2/2)", "公眾查詢網 :8602 — 家長「安心找幼兒園」", 9)
    txt(s, int(0.85 * IN), int(1.42 * IN), int(11.6 * IN), int(0.55 * IN),
        [[("使用者：", 12, True, RISK_GREEN), ("一般民眾 / 家長（無需登入）　｜　", 12, False, INK_MUTED),
          ("刻意與後台分離：", 12, True, INK),
          ("資料層移除所有風險分數欄位，家長端根本拿不到 risk_total", 12, False, INK_MUTED)]],
        line_spacing=1.15)

    # 左：家長三大訴求
    x0 = int(0.85 * IN)
    lw = int(5.6 * IN)
    needs = [("距離近", "輸入地址 → Nominatim 定位 → 找出鄰近幼兒園、地圖標點", "map"),
             ("看得懂學費", "公開收費揭露，透明比較", "money"),
             ("安全放心", "公開評鑑、裁罰狀態、輿情觀測（平靜期/有討論/值得留意）", "shield")]
    yy = int(2.15 * IN)
    for t, d, _ in needs:
        card(s, x0, yy, lw, int(1.15 * IN), fill=WHITE)
        rect(s, x0, yy, int(0.09 * IN), int(1.15 * IN), RISK_GREEN)
        txt(s, x0 + int(0.3 * IN), yy + int(0.16 * IN), lw - int(0.5 * IN),
            int(0.4 * IN), [(t, 15, True, INK)])
        txt(s, x0 + int(0.3 * IN), yy + int(0.6 * IN), lw - int(0.5 * IN),
            int(0.5 * IN), [(d, 11, False, INK_MUTED)], line_spacing=1.1)
        yy += int(1.3 * IN)

    # 右：設計原則 + 分離說明
    rx = x0 + lw + int(0.4 * IN)
    rw = int(11.6 * IN) - lw - int(0.4 * IN)
    card(s, rx, int(2.15 * IN), rw, int(3.75 * IN), fill=SURFACE)
    rect(s, rx, int(2.15 * IN), rw, int(0.55 * IN), RISK_GREEN, rounded=True)
    rect(s, rx, int(2.5 * IN), rw, int(0.2 * IN), RISK_GREEN)
    txt(s, rx + int(0.3 * IN), int(2.22 * IN), rw - int(0.6 * IN),
        int(0.42 * IN), [("設計原則：溫暖、誠實、不恐慌", 15, True, WHITE)],
        anchor=MSO_ANCHOR.MIDDLE)
    points = [
        "不使用後台「高/中/低風險」字眼，改用描述性「近期關注度」並附依據",
        "地圖用柔和暖色階（非警報紅），避免造成家長不必要恐慌",
        "輿情觀測即時爬公開新聞 RSS，每則附原文連結與發布時間，可追溯",
        "地址定位用 OpenStreetMap Nominatim（免金鑰，不用 Google Maps）",
        "資料層 authorize_dataframe(df, ROLE_PARENT) 二次防禦，權限在資料層擋",
    ]
    py = int(2.95 * IN)
    for p in points:
        rect(s, rx + int(0.32 * IN), py + int(0.08 * IN), int(0.12 * IN),
             int(0.12 * IN), RISK_GREEN)
        txt(s, rx + int(0.55 * IN), py, rw - int(0.85 * IN), int(0.6 * IN),
            [(p, 11.5, False, INK)], line_spacing=1.1)
        py += int(0.58 * IN)


# ================================================================ SLIDE 10 收尾
def s_close():
    s = slide()
    rect(s, 0, 0, SW, SH, PRIMARY_DK)
    rect(s, 0, 0, int(0.16 * IN), SH, PRIMARY)
    txt(s, int(0.85 * IN), int(0.85 * IN), int(11.5 * IN), int(0.5 * IN),
        [("為什麼是我們", 15, True, RGBColor(0x9F, 0xBE, 0xDC))])
    txt(s, int(0.85 * IN), int(1.3 * IN), int(11.5 * IN), int(0.7 * IN),
        [("鑑識會計護城河 × 白盒可解釋 × 誠實可規模化", 30, True, WHITE)])

    items = [
        ("護城河", "鑑識會計多層方法，每層有文獻佐證。90% 隊伍不會講「鑑識會計」四個字。"),
        ("政府敢用", "白盒子分數可攤開拆解，責任 AI 全程「風險 ≠ 違法」附證據鏈，禁得起追問。"),
        ("真的能跑", "雙 app live on AWS EC2、61 間實測評分、AUC 0.906，Demo 當場可操作。"),
        ("誠實可擴", "抽樣深度 61 間 + 全市涵蓋 1,114 間，資料來源三態標示，架構可規模化至全市。"),
    ]
    x0 = int(0.85 * IN)
    gap = int(0.3 * IN)
    cw = (int(11.5 * IN) - gap) // 2
    ch = int(1.55 * IN)
    y0 = int(2.5 * IN)
    for i, (t, d) in enumerate(items):
        col = i % 2
        row = i // 2
        x = x0 + col * (cw + gap)
        y = y0 + row * (ch + int(0.25 * IN))
        rect(s, x, y, cw, ch, RGBColor(0x24, 0x42, 0x60), rounded=True)
        rect(s, x, y, int(0.09 * IN), ch, RGBColor(0xFF, 0xD9, 0x7A))
        txt(s, x + int(0.3 * IN), y + int(0.2 * IN), cw - int(0.5 * IN),
            int(0.4 * IN), [(t, 16, True, RGBColor(0xFF, 0xD9, 0x7A))])
        txt(s, x + int(0.3 * IN), y + int(0.68 * IN), cw - int(0.5 * IN),
            int(0.8 * IN), [(d, 12, False, RGBColor(0xDD, 0xE7, 0xF1))],
            line_spacing=1.2)

    txt(s, int(0.85 * IN), int(6.35 * IN), int(11.5 * IN), int(0.6 * IN),
        [[("小小守護員 Smart Watchdog", 16, True, WHITE),
          ("　｜　讓事後被動稽查，提前為事前主動示警。", 14, False,
           RGBColor(0xB9, 0xCA, 0xDD))]])


# ---------------------------------------------------------------- build
def main():
    s_cover()
    s_pain()
    s_solution()
    s_forensic()
    s_whitebox()
    s_data()
    s_aws()
    s_ui_gov()
    s_ui_public()
    s_close()
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    print(f"[OK] 已產出 {OUT}（{len(prs.slides._sldIdLst)} 頁）")


if __name__ == "__main__":
    main()
