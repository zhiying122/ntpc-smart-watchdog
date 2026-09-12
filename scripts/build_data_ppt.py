# -*- coding: utf-8 -*-
"""
小小守護員 Smart Watchdog — 補強簡報：數據及資料運用（靜態 / 動態 API）
=========================================================================
聚焦評審「資料應用性 30%」。內容全依實際程式碼：
  - src/live_source.py（動態連接器：HTTPS GET + 逾時 + 三層備援）
  - docs/data-integration-plan.md（資料介接與備援計畫書）
  - src/data_source_matrix.py（誠實三態資料清冊）
  - 靜態：公校決算書 PDF 抽取、非營利園 OCR/Bedrock

用法：python scripts/build_data_ppt.py
輸出：reports/小小守護員_資料運用補強簡報.pptx
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "reports", "小小守護員_資料運用補強簡報.pptx")

EMU = 914400
IN = EMU
SW = int(13.333 * EMU)
SH = int(7.5 * EMU)

INK = RGBColor(0x1A, 0x1F, 0x2B)
INK_MUTED = RGBColor(0x5B, 0x63, 0x72)
BG = RGBColor(0xFF, 0xFF, 0xFF)
SURFACE = RGBColor(0xF4, 0xF6, 0xF9)
SURFACE2 = RGBColor(0xEA, 0xEE, 0xF3)
BORDER = RGBColor(0xD5, 0xDB, 0xE3)
PRIMARY = RGBColor(0x1F, 0x4E, 0x79)
PRIMARY_DK = RGBColor(0x14, 0x33, 0x50)
ACCENT = RGBColor(0x2C, 0x6E, 0x9B)
RISK_RED = RGBColor(0xD6, 0x45, 0x45)
RISK_ORANGE = RGBColor(0xE0, 0x7B, 0x2B)
RISK_YELLOW = RGBColor(0xE0, 0xA9, 0x3B)
RISK_GREEN = RGBColor(0x3F, 0x8F, 0x5B)
AWS_ORANGE = RGBColor(0xEC, 0x91, 0x2D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD = RGBColor(0xFF, 0xD9, 0x7A)

FONT = "Microsoft JhengHei"

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, color, line=None, line_w=None, rounded=False):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    if color is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = color
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w or 1)
    shp.shadow.inherit = False
    return shp


def arrow(s, x, y, w, h, color):
    shp = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Emu(int(x)), Emu(int(y)),
                             Emu(int(w)), Emu(int(h)))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def down_arrow(s, x, y, w, h, color):
    shp = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Emu(int(x)), Emu(int(y)),
                             Emu(int(w)), Emu(int(h)))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space_after=4, line_spacing=1.0, wrap=True):
    tb = s.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)),
                              Emu(int(h)))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    if runs and isinstance(runs[0], tuple):
        paras = [runs]
    else:
        paras = runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (t, size, bold, color) in para:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = FONT
            r.font.color.rgb = color
    return tb


def card(s, x, y, w, h, fill=SURFACE, line=BORDER):
    return rect(s, x, y, w, h, fill, line=line, line_w=1, rounded=True)


def chip(s, x, y, w, h, label, fill, fg=WHITE, size=11, bold=True):
    c = rect(s, x, y, w, h, fill, rounded=True)
    tf = c.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(3)
    tf.margin_right = Pt(3)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = FONT
    r.font.color.rgb = fg
    return c


def header(s, kicker, title, idx=None):
    rect(s, 0, 0, SW, int(1.35 * IN), BG)
    rect(s, int(0.6 * IN), int(0.42 * IN), int(0.09 * IN), int(0.62 * IN),
         PRIMARY)
    txt(s, int(0.85 * IN), int(0.38 * IN), int(10 * IN), int(0.3 * IN),
        [(kicker, 12, True, ACCENT)])
    txt(s, int(0.85 * IN), int(0.62 * IN), int(11.2 * IN), int(0.62 * IN),
        [(title, 26, True, INK)])
    rect(s, int(0.85 * IN), int(1.24 * IN), int(11.6 * IN), Emu(12700), BORDER)
    if idx is not None:
        txt(s, int(11.9 * IN), int(0.5 * IN), int(0.9 * IN), int(0.3 * IN),
            [(f"{idx:02d}", 12, True, BORDER)], align=PP_ALIGN.RIGHT)


# ============================================================ 1 封面
def s_cover():
    s = slide()
    rect(s, 0, 0, SW, SH, PRIMARY_DK)
    rect(s, 0, 0, int(0.16 * IN), SH, PRIMARY)
    chip(s, int(0.85 * IN), int(1.0 * IN), int(4.2 * IN), int(0.42 * IN),
         "補強簡報 ｜ 資料應用性 (評分 30%)", RGBColor(0x2A, 0x4A, 0x6B),
         fg=RGBColor(0xCF, 0xE0, 0xF0), size=13)
    txt(s, int(0.85 * IN), int(2.1 * IN), int(11.5 * IN), int(1.2 * IN),
        [("數據及資料運用", 48, True, WHITE)])
    txt(s, int(0.85 * IN), int(3.3 * IN), int(11.5 * IN), int(0.7 * IN),
        [[("靜態一手資料 ", 22, True, GOLD),
          ("×", 22, False, WHITE),
          (" 動態官方 API 整合 ", 22, True, GOLD),
          ("×", 22, False, WHITE),
          (" 三層備援韌性", 22, True, GOLD)]])
    txt(s, int(0.85 * IN), int(4.5 * IN), int(11.5 * IN), int(1.4 * IN),
        [[("回應題目痛點 1：", 15, True, RGBColor(0xC7, 0xD6, 0xE6)),
          ("「資料分散、整合困難、缺乏即時整合機制」", 15, False,
           RGBColor(0xC7, 0xD6, 0xE6))],
         [("我們把系統從「靜態一次性資料」升級為「動態串接官方開放資料 + 本地快取 + 離線備援」的整合平台。",
           14, False, RGBColor(0xA9, 0xC0, 0xD8))]],
        line_spacing=1.3, space_after=8)
    txt(s, int(0.85 * IN), int(6.4 * IN), int(11.5 * IN), int(0.4 * IN),
        [("小小守護員 Smart Watchdog ｜ 2026 新北市 AI 智慧城市黑客松 · 教育局組",
          11, False, RGBColor(0x6E, 0x8A, 0xA6))])


# ============================================================ 2 資料全景
def s_landscape():
    s = slide()
    header(s, "DATA LANDSCAPE ｜ 資料全景", "五大類資料，橫跨財務、監理、公眾三面向", 2)
    txt(s, int(0.85 * IN), int(1.42 * IN), int(11.6 * IN), int(0.4 * IN),
        [("題目點名的基本資料、評鑑、裁罰、收費、決算財報，我們全數整合，並補上社群輿情面向。",
          13, False, INK_MUTED)], line_spacing=1.1)

    cats = [
        ("財務決算", "公校決算書 112–114 年度\n非營利園財報 110–113", "鑑識會計核心輸入", RISK_RED),
        ("裁罰紀錄", "全國教保資訊網\n(含日期/法條/罰鍰/文號)", "風險標籤與驗證基準", RISK_ORANGE),
        ("評鑑結果", "幼兒園評鑑等第\n優/良/乙/待改進", "評鑑分項輸入", RISK_YELLOW),
        ("基本 / 收費", "園名/行政區/地址/座標\n核定人數/月收費", "定位 + 交叉勾稽 + 家長端", ACCENT),
        ("網路輿情", "公開新聞 / 社群 / 評論\n(五層來源分級)", "微弱訊號早期預警", RISK_GREEN),
    ]
    x0 = int(0.85 * IN)
    gap = int(0.22 * IN)
    cw = (int(11.6 * IN) - 4 * gap) // 5
    y0 = int(2.15 * IN)
    ch = int(3.0 * IN)
    for i, (name, src, use, color) in enumerate(cats):
        x = x0 + i * (cw + gap)
        card(s, x, y0, cw, ch, fill=WHITE)
        rect(s, x, y0, cw, int(0.7 * IN), color, rounded=True)
        rect(s, x, y0 + int(0.45 * IN), cw, int(0.25 * IN), color)
        txt(s, x + int(0.12 * IN), y0 + int(0.08 * IN), cw - int(0.24 * IN),
            int(0.55 * IN), [(name, 15, True, WHITE)],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        txt(s, x + int(0.16 * IN), y0 + int(0.85 * IN), cw - int(0.3 * IN),
            int(1.3 * IN), [(src, 11, False, INK)], line_spacing=1.2)
        rect(s, x + int(0.16 * IN), y0 + ch - int(0.85 * IN),
             cw - int(0.32 * IN), Emu(12700), BORDER)
        txt(s, x + int(0.16 * IN), y0 + ch - int(0.72 * IN),
            cw - int(0.3 * IN), int(0.6 * IN),
            [[("用途：", 9.5, True, color), (use, 9.5, False, INK_MUTED)]],
            line_spacing=1.05)

    txt(s, x0, int(5.5 * IN), int(11.6 * IN), int(0.5 * IN),
        [[("交叉勾稽：", 13, True, PRIMARY),
          ("財報 × 收費 × 核定人數 → 算「每名幼兒單位成本」；裁罰 × 風險分 → 驗證高風險園是否確實較常被裁罰（AUC 0.906）。",
           12, False, INK_MUTED)]], line_spacing=1.15)


# ============================================================ 3 靜態 vs 動態雙軌
def s_static_dynamic():
    s = slide()
    header(s, "TWO-TRACK ｜ 靜態 × 動態雙軌設計", "一手精抽 + 官方即時，各司其職", 3)

    x0 = int(0.85 * IN)
    gap = int(0.4 * IN)
    cw = (int(11.6 * IN) - gap) // 2
    y0 = int(1.65 * IN)
    ch = int(4.65 * IN)

    # 左：靜態
    card(s, x0, y0, cw, ch, fill=WHITE)
    rect(s, x0, y0, cw, int(0.75 * IN), PRIMARY_DK, rounded=True)
    rect(s, x0, y0 + int(0.5 * IN), cw, int(0.25 * IN), PRIMARY_DK)
    txt(s, x0 + int(0.3 * IN), y0 + int(0.1 * IN), cw - int(0.6 * IN),
        int(0.55 * IN),
        [[("靜態資料 ", 18, True, WHITE), ("STATIC", 13, True, GOLD)]],
        anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x0 + int(0.3 * IN), y0 + int(0.9 * IN), cw - int(0.6 * IN),
        int(0.4 * IN), [("一手官方文件、深度精準抽取", 13, True, INK)])
    static_items = [
        ("公校決算書 PDF", "地方教育發展基金決算書第5冊，PyMuPDF 文字層抽取，含逐筆明細供班佛檢定"),
        ("非營利園掃描財報", "無文字層 → Tesseract 中文 OCR + Bedrock Claude 多模態視覺抽取，自我勾稽把關"),
        ("地理座標（衍生）", "Nominatim geocoding 產生 lat/lng，供地圖標記"),
    ]
    yy = y0 + int(1.4 * IN)
    for t, d in static_items:
        rect(s, x0 + int(0.32 * IN), yy + int(0.06 * IN), int(0.13 * IN),
             int(0.13 * IN), PRIMARY)
        txt(s, x0 + int(0.55 * IN), yy, cw - int(0.85 * IN), int(0.35 * IN),
            [(t, 13, True, INK)])
        txt(s, x0 + int(0.55 * IN), yy + int(0.34 * IN), cw - int(0.85 * IN),
            int(0.7 * IN), [(d, 10.5, False, INK_MUTED)], line_spacing=1.1)
        yy += int(1.02 * IN)
    chip(s, x0 + int(0.32 * IN), y0 + ch - int(0.62 * IN), cw - int(0.64 * IN),
         int(0.42 * IN), "特性：正確性高、可追溯至來源頁，適合精準鑑識",
         SURFACE2, fg=INK, size=11)

    # 右：動態
    rx = x0 + cw + gap
    card(s, rx, y0, cw, ch, fill=WHITE)
    rect(s, rx, y0, cw, int(0.75 * IN), RISK_GREEN, rounded=True)
    rect(s, rx, y0 + int(0.5 * IN), cw, int(0.25 * IN), RISK_GREEN)
    txt(s, rx + int(0.3 * IN), y0 + int(0.1 * IN), cw - int(0.6 * IN),
        int(0.55 * IN),
        [[("動態資料 ", 18, True, WHITE), ("LIVE API", 13, True, GOLD)]],
        anchor=MSO_ANCHOR.MIDDLE)
    txt(s, rx + int(0.3 * IN), y0 + int(0.9 * IN), cw - int(0.6 * IN),
        int(0.4 * IN), [("官方開放資料即時串接、隨更新重評估", 13, True, INK)])
    dyn_items = [
        ("幼兒園基本資料 / 收費", "preschools.json（GeoJSON，含座標），HTTPS GET 免金鑰"),
        ("全國裁罰紀錄", "punish_all.json，含日期/法條/罰鍰/文號，即時解析正規化"),
        ("連接器 src/live_source.py", "load_live_dataset() 一鍵同步；UI「同步最新資料」按鈕觸發"),
    ]
    yy = y0 + int(1.4 * IN)
    for t, d in dyn_items:
        rect(s, rx + int(0.32 * IN), yy + int(0.06 * IN), int(0.13 * IN),
             int(0.13 * IN), RISK_GREEN)
        txt(s, rx + int(0.55 * IN), yy, cw - int(0.85 * IN), int(0.35 * IN),
            [(t, 13, True, INK)])
        txt(s, rx + int(0.55 * IN), yy + int(0.34 * IN), cw - int(0.85 * IN),
            int(0.7 * IN), [(d, 10.5, False, INK_MUTED)], line_spacing=1.1)
        yy += int(1.02 * IN)
    chip(s, rx + int(0.32 * IN), y0 + ch - int(0.62 * IN), cw - int(0.64 * IN),
         int(0.42 * IN), "特性：涵蓋廣、時效新，排程每日更新即成即時預警系統",
         SURFACE2, fg=INK, size=11)


# ============================================================ 4 動態連接器 + 三層備援
def s_resilience():
    s = slide()
    header(s, "RELIABILITY ｜ 動態連接器與韌性設計", "政府級韌性：外部源掛了，系統也不倒", 4)
    txt(s, int(0.85 * IN), int(1.42 * IN), int(11.6 * IN), int(0.4 * IN),
        [[("src/live_source.py 的 ", 12, False, INK_MUTED),
          ("fetch_json()", 12, True, PRIMARY),
          (" 讀取優先序：即時網路 → 本地快取 → 版控快照，任一層可用即回傳。",
           12, False, INK_MUTED)]], line_spacing=1.1)

    # 流程主軸
    x0 = int(0.85 * IN)
    y = int(2.1 * IN)
    bw = int(2.45 * IN)
    bh = int(0.95 * IN)
    gapx = int(0.55 * IN)

    def flowbox(x, title, sub, color, fg=WHITE):
        rect(s, x, y, bw, bh, color, rounded=True)
        txt(s, x + int(0.1 * IN), y + int(0.12 * IN), bw - int(0.2 * IN),
            int(0.4 * IN), [(title, 13, True, fg)], align=PP_ALIGN.CENTER)
        txt(s, x + int(0.1 * IN), y + int(0.5 * IN), bw - int(0.2 * IN),
            int(0.4 * IN), [(sub, 9.5, False, fg)], align=PP_ALIGN.CENTER,
            line_spacing=1.0)

    flowbox(x0, "官方開放資料源", "教育部 / 全國教保資訊網", RGBColor(0x5B, 0x63, 0x72))
    arrow(s, x0 + bw + int(0.08 * IN), y + int(0.32 * IN), gapx - int(0.16 * IN),
          int(0.3 * IN), AWS_ORANGE)
    flowbox(x0 + (bw + gapx), "HTTPS GET", "urllib · timeout 30s · 免金鑰",
            AWS_ORANGE)
    arrow(s, x0 + 2 * bw + gapx + int(0.08 * IN), y + int(0.32 * IN),
          gapx - int(0.16 * IN), int(0.3 * IN), PRIMARY)
    flowbox(x0 + 2 * (bw + gapx), "連接器解析正規化", "parse_preschools / parse_penalties",
            PRIMARY)
    arrow(s, x0 + 3 * bw + 2 * gapx + int(0.08 * IN), y + int(0.32 * IN),
          gapx - int(0.16 * IN), int(0.3 * IN), RISK_GREEN)
    flowbox(x0 + 3 * (bw + gapx), "風險引擎 + UI", "契約檔 → 三入口介面", RISK_GREEN)

    # 三層備援
    ty = int(3.55 * IN)
    txt(s, x0, ty, int(11.6 * IN), int(0.35 * IN),
        [("三層備援（Fallback Chain）— FetchResult 誠實回報資料新鮮度", 14, True, INK)])
    layers = [
        ("① 即時網路", "抓取成功 → 寫本地快取", "is_live = True", RISK_GREEN,
         "資料最新，UI 標「即時串接」"),
        ("② 本地快取 data/cache", "抓取失敗 → 退回最近快取", "used_fallback = True", RISK_YELLOW,
         "UI 標「離線快取備援」+ 抓取時間"),
        ("③ 版控快照 data/snapshots", "GitHub Actions 每日更新進版控", "線上部署也讀得到", ACCENT,
         "無網路仍有近一天內新鮮資料"),
    ]
    lw = (int(11.6 * IN) - 2 * int(0.3 * IN)) // 3
    ly = int(3.95 * IN)
    lh = int(2.35 * IN)
    for i, (t, act, flag, color, note) in enumerate(layers):
        x = x0 + i * (lw + int(0.3 * IN))
        card(s, x, ly, lw, lh, fill=WHITE)
        rect(s, x, ly, lw, int(0.6 * IN), color, rounded=True)
        rect(s, x, ly + int(0.35 * IN), lw, int(0.25 * IN), color)
        txt(s, x + int(0.2 * IN), ly + int(0.08 * IN), lw - int(0.4 * IN),
            int(0.45 * IN), [(t, 14, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
        txt(s, x + int(0.25 * IN), ly + int(0.78 * IN), lw - int(0.45 * IN),
            int(0.6 * IN), [(act, 12, True, INK)], line_spacing=1.1)
        chip(s, x + int(0.25 * IN), ly + int(1.42 * IN), lw - int(0.5 * IN),
             int(0.38 * IN), flag, SURFACE2, fg=PRIMARY, size=10.5)
        txt(s, x + int(0.25 * IN), ly + int(1.9 * IN), lw - int(0.45 * IN),
            int(0.4 * IN), [(note, 10, False, INK_MUTED)], line_spacing=1.05)


# ============================================================ 5 資料治理與誠實揭露
def s_governance():
    s = slide()
    header(s, "GOVERNANCE ｜ 資料治理與誠實揭露", "政府敢用的前提：資料誠實、可追溯、不虛構", 5)

    x0 = int(0.85 * IN)
    # 左：三態資料清冊說明
    lw = int(5.5 * IN)
    card(s, x0, int(1.6 * IN), lw, int(2.5 * IN), fill=WHITE)
    txt(s, x0 + int(0.3 * IN), int(1.78 * IN), lw - int(0.6 * IN),
        int(0.4 * IN),
        [("資料來源矩陣三態標記", 15, True, INK)])
    txt(s, x0 + int(0.3 * IN), int(2.18 * IN), lw - int(0.6 * IN),
        int(0.35 * IN),
        [("src/data_source_matrix.py（誠實清冊 R16）", 10.5, False, INK_MUTED)])
    tri = [("已確認 CONFIRMED", "實際持有、已驗證可用", RISK_GREEN),
           ("未確認 UNCONFIRMED", "尚未取得，誠實標示不宣稱可用", RISK_YELLOW),
           ("不適用 N/A", "該屬性對此資料源不適用", INK_MUTED)]
    yy = int(2.65 * IN)
    for name, d, color in tri:
        chip(s, x0 + int(0.3 * IN), yy, int(2.0 * IN), int(0.38 * IN), name,
             color, size=11)
        txt(s, x0 + int(2.45 * IN), yy + int(0.02 * IN), int(2.9 * IN),
            int(0.4 * IN), [(d, 10.5, False, INK)], anchor=MSO_ANCHOR.MIDDLE,
            line_spacing=1.0)
        yy += int(0.48 * IN)
    txt(s, x0 + int(0.3 * IN), int(3.55 * IN), lw - int(0.6 * IN),
        int(0.5 * IN),
        [("原則：只列實際存在的資料集，絕不虛構；官方全量 API 未取得者一律標「未確認」。",
          10.5, False, INK_MUTED)], line_spacing=1.15)

    # 右：誠實標註要點
    rx = x0 + lw + int(0.4 * IN)
    rw = int(11.6 * IN) - lw - int(0.4 * IN)
    card(s, rx, int(1.6 * IN), rw, int(2.5 * IN), fill=SURFACE)
    txt(s, rx + int(0.3 * IN), int(1.78 * IN), rw - int(0.6 * IN),
        int(0.4 * IN), [("來源標註與資料治理", 15, True, INK)])
    pts = [
        "動態 JSON 為 g0v 開源專案（江明宗）整理備份，原始來自全國教保資訊網",
        "系統一律標註來源 + 抓取時間，供時效判斷與追溯",
        "二手資料明示「需抽查」，上架前核對 5–10 筆與官方一致",
        "正式部署改接官方 API：data.gov.tw 6086 / 新北市 OpenAPI",
    ]
    py = int(2.25 * IN)
    for p in pts:
        rect(s, rx + int(0.32 * IN), py + int(0.07 * IN), int(0.12 * IN),
             int(0.12 * IN), PRIMARY)
        txt(s, rx + int(0.55 * IN), py, rw - int(0.85 * IN), int(0.5 * IN),
            [(p, 11.5, False, INK)], line_spacing=1.1)
        py += int(0.46 * IN)

    # 底：責任 AI + 資安
    y = int(4.4 * IN)
    bottom = [
        ("責任 AI", "風險 ≠ 違法；低可信度資料不得單獨推高風險；每則風險陳述附聲明", RISK_GREEN),
        ("實體解析", "裁罰以個人姓名為鍵，不確定比對標「待人工確認」不強制合併", ACCENT),
        ("資安合規", "唯讀 HTTPS GET 不外傳資料；金鑰放 .env，check_secrets.py 上傳前掃描", PRIMARY),
    ]
    bw = (int(11.6 * IN) - 2 * int(0.3 * IN)) // 3
    for i, (t, d, color) in enumerate(bottom):
        x = x0 + i * (bw + int(0.3 * IN))
        card(s, x, y, bw, int(1.85 * IN), fill=WHITE)
        rect(s, x, y, bw, int(0.07 * IN), color)
        txt(s, x + int(0.25 * IN), y + int(0.2 * IN), bw - int(0.45 * IN),
            int(0.4 * IN), [(t, 14, True, color)])
        txt(s, x + int(0.25 * IN), y + int(0.65 * IN), bw - int(0.45 * IN),
            int(1.1 * IN), [(d, 11, False, INK_MUTED)], line_spacing=1.2)


# ============================================================ 6 資料→價值
def s_pipeline():
    s = slide()
    header(s, "DATA TO VALUE ｜ 資料如何變成稽查決策", "從原始資料到可解釋風險分，一條龍", 6)

    x0 = int(0.85 * IN)
    steps = [
        ("原始資料", "決算 PDF / 裁罰 JSON /\n收費 / 評鑑 / 輿情", RGBColor(0x5B, 0x63, 0x72)),
        ("抽取正規化", "PDF 文字層 / OCR /\nBedrock / API 解析", AWS_ORANGE),
        ("鑑識會計指標", "班佛 / Beneish / Altman /\n孤立森林 / 勾稽", PRIMARY),
        ("白盒風險分", "財務50 + 裁罰34 +\n評鑑16，可攤開", RISK_ORANGE),
        ("稽查決策", "排名 / 地圖 / 雷達圖 /\nAI建議 / 派工", RISK_GREEN),
    ]
    n = len(steps)
    bw = int(2.0 * IN)
    gapx = (int(11.6 * IN) - n * bw) // (n - 1)
    y = int(1.9 * IN)
    bh = int(1.5 * IN)
    for i, (t, d, color) in enumerate(steps):
        x = x0 + i * (bw + gapx)
        rect(s, x, y, bw, bh, WHITE, line=color, line_w=2, rounded=True)
        rect(s, x, y, bw, int(0.5 * IN), color, rounded=True)
        rect(s, x, y + int(0.3 * IN), bw, int(0.2 * IN), color)
        txt(s, x + int(0.1 * IN), y + int(0.06 * IN), bw - int(0.2 * IN),
            int(0.42 * IN), [(t, 13, True, WHITE)], align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
        txt(s, x + int(0.12 * IN), y + int(0.62 * IN), bw - int(0.24 * IN),
            int(0.85 * IN), [(d, 10, False, INK)], align=PP_ALIGN.CENTER,
            line_spacing=1.15)
        if i < n - 1:
            arrow(s, x + bw + int(0.06 * IN), y + int(0.55 * IN),
                  gapx - int(0.12 * IN), int(0.32 * IN), color)

    # 實際成果數字帶
    txt(s, x0, int(3.75 * IN), int(11.6 * IN), int(0.35 * IN),
        [("實際運算成果（真實資料，非模擬）", 14, True, INK)])
    kpis = [("61 間", "已完成鑑識會計風險評分"),
            ("1,114 間", "全市立案機構納管基本資料"),
            ("2,838 筆", "真實裁罰紀錄動態串接"),
            ("AUC 0.906", "以官方裁罰驗證模型鑑別力")]
    kw = (int(11.6 * IN) - 3 * int(0.28 * IN)) // 4
    ky = int(4.2 * IN)
    for i, (v, d) in enumerate(kpis):
        x = x0 + i * (kw + int(0.28 * IN))
        card(s, x, ky, kw, int(1.35 * IN), fill=SURFACE)
        rect(s, x, ky, kw, int(0.07 * IN), PRIMARY)
        txt(s, x, ky + int(0.22 * IN), kw, int(0.5 * IN),
            [(v, 25, True, PRIMARY)], align=PP_ALIGN.CENTER)
        txt(s, x + int(0.12 * IN), ky + int(0.82 * IN), kw - int(0.24 * IN),
            int(0.45 * IN), [(d, 10.5, False, INK_MUTED)],
            align=PP_ALIGN.CENTER, line_spacing=1.05)

    # 收斂句
    rect(s, x0, int(5.85 * IN), int(11.6 * IN), int(0.85 * IN),
         PRIMARY_DK, rounded=True)
    txt(s, x0 + int(0.4 * IN), int(5.9 * IN), int(11.0 * IN), int(0.75 * IN),
        [[("一句話：", 14, True, GOLD),
          ("我們不只「有資料」——靜態一手精抽保證正確、動態官方 API 保證新鮮、三層備援保證不中斷，", 13, False, WHITE)],
         [("再經鑑識會計轉成可解釋風險分，讓每一筆公開資料都變成稽查員的決策依據。", 13, False, WHITE)]],
        anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.2, space_after=2)


def main():
    s_cover()
    s_landscape()
    s_static_dynamic()
    s_resilience()
    s_governance()
    s_pipeline()
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    print(f"[OK] 已產出 {OUT}（{len(prs.slides._sldIdLst)} 頁）")


if __name__ == "__main__":
    main()
